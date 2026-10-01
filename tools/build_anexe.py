#!/usr/bin/env python3
"""Genereaza sabloanele Word (anexe/Anexa_NN.docx) pentru anexele Ordinului
ANCPI nr. 600/2023, pornind de la textul oficial extras in
tools/ordin600_anexe.json (vezi tools/extract_anexe.py).

    python3 tools/build_anexe.py            # scrie anexe/*.docx
    python3 tools/build_anexe.py --raport   # afiseaza campurile detectate

Spatiile punctate din formulare sunt inlocuite cu campuri de forma
{{CHEIE|.......}}: la completare (anexe/TopoKitDocx.ps1) campul primeste
valoarea din desen/dialog sau, daca valoarea lipseste, punctele originale,
astfel incat formularul ramane completabil de mana.

Campurile sunt detectate automat dupa context (cine este descris: solicitantul,
imobilul sau persoana autorizata; eticheta dinaintea punctelor: "str.", "nr.",
"CNP" etc.), plus reguli explicite pentru anexele cu structura speciala.
"""
import json, re, sys, pathlib
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

ROOT = pathlib.Path(__file__).resolve().parent.parent
SRC = ROOT / "tools" / "ordin600_anexe.json"
OUT = ROOT / "anexe"

# --------------------------------------------------------------------------
# Titlurile anexelor (pentru nume de fisier si catalog)
# --------------------------------------------------------------------------
TITLURI = {
    1: "Cerere eliberare copii din arhiva",
    2: "Instiintare si comunicare incheiere (OCPI)",
    3: "Comunicarea incheierii (OCPI)",
    4: "Incheieri (OCPI)",
    5: "Cerere de inscriere",
    6: "Cerere extras de carte funciara pentru informare",
    7: "Cerere certificat de sarcini",
    8: "Cerere identificare numar CF dupa nume proprietar",
    9: "Certificat identificare numar CF (OCPI)",
    10: "Cerere certificat imobil neinscris",
    11: "Certificat imobil neinscris (OCPI)",
    12: "Cerere de informatii",
    13: "Cerere de receptie si inscriere",
    14: "Declaratie proprietar",
    15: "Memoriu tehnic",
    19: "Foaia colectiva",
    20: "Proces-verbal de vecinatate",
    21: "Invitatie lamuriri (OCPI)",
    22: "Cerere de anulare numar cadastral",
    23: "Cerere aviz de incepere-receptie (CNC)",
    24: "Aviz de incepere a lucrarii (CNC)",
    25: "Proces-verbal de receptie (CNC-OCPI)",
    26: "Adeverinta primarie (posesie)",
    27: "Certificat registrul agricol (primarie)",
    28: "Certificat atestare edificare cu AC (primarie)",
    29: "Certificat atestare edificare fara AC (primarie)",
    30: "Cerere de informatii plan parcelar",
    32: "Tabel parcelar",
    33: "Proces-verbal receptie plan parcelar (OCPI)",
    34: "Cerere consultare registru proprietari",
    35: "Certificat de sarcini (OCPI)",
}
# Anexele 16, 17, 18 si 31 sunt piese desenate: se genereaza in AutoCAD.
GRAFICE = {16, 17, 18, 31}
# Formularele lucrate pe orizontala (tabele late)
LANDSCAPE = {19, 32}

PH = re.compile(r"\{\{[^{}]*\}\}")
DOTS = re.compile(r"(?:\.{2,}|…+)(?:[ ]?/[ ]?(?:20)?(?:\.{2,}|…+)|[ ]?(?:\.{2,}|…+))*")

def ph(key, dots):
    return "{{%s|%s}}" % (key, dots)

# --------------------------------------------------------------------------
# Detectia automata a campurilor
# --------------------------------------------------------------------------
TRIG_P = ["subsemnatul", "subsemnata", "subscrisa", "domnul ", "proprietatea domnului",
          "deţinătorul construcţiei", "cererii depuse de către", "proprietatea numitului"]
TRIG_P2 = ["şi doamna"]
TRIG_IMOB = ["imobilul situat", "privind imobilul", "imobilul este identificat", "al imobilului",
             "imobilului situat", "imobilul: uat", "cu privire la imobilul", "pe imobilul situat",
             "este situat la adresa", "intravilan, localitatea", "extravilan, tarla", "situat în str",
             "împreună cu construcţia, situat", "imobilul format", "imobilul ", "terenul în suprafaţă"]
TRIG_EXEC = ["împuternicesc", "persoanei autorizate", "persoana fizică/juridică autorizată",
             "persoana autorizată"]
TRIG_NONE = ["s-a achitat", "nr. de înregistrare", "tariful", "proprietarul anterior",
             "pentru intervalul", "litigiu", "solicit comunicarea", "înregistrată cu",
             "autorizaţiei de construire", "proces-verbal de recepţie", "expertiza tehnică",
             "ca urmare a cererii", "prin mandatarul"]
# Anexele emise de OCPI/CNC: se completeaza de institutie, raman necompletate
AUTO_OFF = {2, 3, 4, 9, 11, 21, 24, 25, 33, 35}
# Etichete care descriu intotdeauna imobilul, indiferent de subiectul frazei
IMOB_ALWAYS = {"UAT", "NRCF", "NRCAD", "TARLA", "PARCELA", "NRTOP", "SUPR", "VECINATATI", "CATEG"}

LABELS = [
    (r"(oficiul de cadastru şi publicitate imobiliară|oficiul de cadastru şi publicitate imobiliara)$", "OCPI"),
    (r"(biroul de cadastru şi publicitate imobiliară)$", "BCPI"),
    (r"(subsemnatul\s*\(a\),?|subsemnatul/\(a\)/subscrisa,?|subsemnata/subsemnatul/subscrisa|"
     r"subsemnatul\(a\),?|domnul|şi doamna|numitului/denumirea|numitul|numitului|domnului/doamnei|"
     r"depuse de către|împuternicesc persoana fizică/juridică|persoanei autorizate)$", "NUMEC"),
    (r"(cnp/cui|cnp/cui/cif|cnp)$", "CNP"),
    (r"(domiciliat\(ă\) în localitatea|domiciliat \(ă\) în localitatea|domiciliat\(ă\) în|"
     r"cu domiciliul/sediul în|cu domiciliul în|domiciliat \(ă\) în/cu sediul în localitatea|"
     r"localitatea|domiciliat în|situat în)$", "LOC"),
    (r"(str\.|strada|str\. \.)$", "STR"),
    (r"(bl\.)$", "BL"),
    (r"(sc\.|scara)$", "SC"),
    (r"(et\.|etaj|et)$", "ET"),
    (r"(ap\.|ap)$", "AP"),
    (r"(sectorul/judeţul|judeţul/sectorul|sector/judeţ|judeţul|sector nr\./judeţul)$", "JUD"),
    (r"(cod poştal)$", "CODP"),
    (r"(ci/bi seria|bi/ci seria|bi/ci/paşaport seria|paşaport seria)$", "CISERIA"),
    (r"(categoria/clasa|certificatului de autorizare categorie|autorizaţie categoria/clasa)$", "CAT"),
    (r"(seria)$", "SERIA"),
    (r"(telefon|tel\./fax|tel/fax)$", "TEL"),
    (r"(e-mail:?)$", "EMAIL"),
    (r"(eliberat\(ă\) de|eliberat de)$", "EMITENT"),
    (r"(cartea funciară nr\.|carte funciară nr\.|cartea funciară nr|carte funciară colectivă nr\.)$", "NRCF"),
    (r"(nr\. cadastral|numărul cadastral al imobilului|număr topografic/cadastral|"
     r"nr\. cadastral/topografic|cu nr\. cadastral|cadastral/topografic)$", "NRCAD"),
    (r"(a comunei/oraşului/municipiului|comuna/oraşul/municipiul|a localităţii|uat|uat/zonă|"
     r"primăria comunei/oraşului/municipiului)$", "UAT"),
    (r"(tarla|tarlaua|tarlalei nr\.)$", "TARLA"),
    (r"(parcela|parcelă)$", "PARCELA"),
    (r"(nr\. top)$", "NRTOP"),
    (r"(nr\.|nr|nr \.|nr\. \.)$", "NR"),
    (r"(compus din|în suprafaţă de|rezultate din măsurători de)$", "SUPR"),
    (r"(contractului/angajamentului nr\.)$", "CONTRACT"),
    (r"(data)$", "DATA"),
    (r"(cu următoarele vecinătăţi:?)$", "VECINATATI"),
    (r"(categoria de folosinţă)$", "CATEG"),
]
LABELS = [(re.compile(rx), g) for rx, g in LABELS]

MAP = {
    "P": {"NUMEC": "P1_NUMEC", "CNP": "P1_CNP", "LOC": "P1_LOC", "STR": "P1_STR", "NR": "P1_NR",
          "BL": "P1_BL", "SC": "P1_SC", "ET": "P1_ET", "AP": "P1_AP", "JUD": "P1_JUD",
          "CODP": "P1_CODP", "CISERIA": "P1_CI_SERIA", "CINR": "P1_CI_NR", "TEL": "P1_TEL",
          "EMAIL": "P1_EMAIL", "EMITENT": "P1_CI_EMITENT", "SERIA": "P1_CI_SERIA"},
    "IMOB": {"LOC": "LOCALITATE", "STR": "STRADA", "NR": "NR", "BL": "BLOC", "SC": "SCARA",
             "ET": "ETAJ", "AP": "AP", "JUD": "JUDET", "UAT": "UAT", "NRCF": "NRCF",
             "NRCAD": "NRCAD", "TARLA": "TARLA", "PARCELA": "PARCELA", "NRTOP": "NRTOP",
             "SUPR": "SMAS", "VECINATATI": "VECINATATI", "CATEG": "CATEG_TEXT",
             "NUMEC": "P1_NUMEC", "CNP": "P1_CNP"},
    "EXEC": {"NUMEC": "E_NUMEC", "CNP": "E_CNP", "LOC": "E_LOC", "STR": "E_STR", "NR": "E_NR",
             "BL": "E_BL", "SC": "E_SC", "ET": "E_ET", "AP": "E_AP", "JUD": "E_JUD",
             "CISERIA": "E_CI_SERIA", "CINR": "E_CI_NR", "TEL": "E_TEL", "EMAIL": "E_EMAIL",
             "CAT": "E_CAT", "SERIA": "E_AUT_SERIA", "AUTNR": "E_AUT_NR",
             "EMITENT": "E_AUT_EMITENT", "DATA": "E_AUT_DATA", "CONTRACT": "E_CONTRACT",
             "CONTRACTDATA": "E_CONTRACT_DATA"},
}
MAP["P2"] = {k: v.replace("P1_", "P2_") for k, v in MAP["P"].items()}

class Engine:
    """Masina de stari care atribuie chei spatiilor punctate dintr-o anexa."""

    def __init__(self, nr, exec_subject=False):
        self.nr = nr
        self.exec_subject = exec_subject   # solicitantul este persoana autorizata (anexa 23)
        self.subj = None
        self.prev = None
        self.found = []

    def triggers(self, low):
        best = None
        for lst, subj in ((TRIG_P, "P"), (TRIG_P2, "P2"), (TRIG_IMOB, "IMOB"),
                          (TRIG_EXEC, "EXEC"), (TRIG_NONE, None)):
            for t in lst:
                i = low.rfind(t)
                if i >= 0 and (best is None or i > best[0]):
                    best = (i, subj)
        if best is not None:
            self.subj = best[1]
            if self.exec_subject and self.subj == "P":
                self.subj = "EXEC"

    def classify(self, before):
        low = before.lower().replace("ş", "ş").rstrip(" ,:;(")
        low = re.sub(r"\s+", " ", low)
        g = None
        for rx, gen in LABELS:
            if rx.search(low):
                g = gen
                break
        if g is None:
            return None
        if g == "NR":
            if self.prev in ("CISERIA",):
                g = "CINR"
            elif self.prev == "SERIA" and self.subj == "EXEC":
                g = "AUTNR"
            elif self.prev == "SERIA":
                g = "CINR"
            elif self.prev == "CONTRACT":
                return None
        if g == "DATA":
            if low == "data" or low.endswith("încheiat la data"):
                return "DATA"
            if self.prev == "CONTRACT":
                return "E_CONTRACT_DATA"
            if self.prev == "EMITENT" and self.subj == "EXEC":
                return "E_AUT_DATA"
            return None
        if g == "NUMEC" and "împuternicesc" in low:
            return "E_NUMEC"
        if g in ("OCPI", "BCPI"):
            return g
        if g in IMOB_ALWAYS:
            return MAP["IMOB"][g]
        if g == "JUD" and self.prev == "UAT":
            return "JUDET"
        if self.subj is None:
            return None
        return MAP.get(self.subj, {}).get(g)

    def run(self, text):
        """Inlocuieste punctele cu campuri; textul deja marcat ramane neschimbat."""
        out, pos, ctx = [], 0, ""
        parts = PH.split(text)
        phs = PH.findall(text)
        res = []
        for i, seg in enumerate(parts):
            res.append(self._seg(seg))
            if i < len(phs):
                res.append(phs[i])
        return "".join(res)

    def _seg(self, seg):
        out, pos = [], 0
        for m in DOTS.finditer(seg):
            before = seg[pos:m.start()]
            self.triggers(before.lower())
            label_ctx = (seg[max(0, m.start() - 60):m.start()])
            key = self.classify(label_ctx)
            g_prev = None
            low = re.sub(r"\s+", " ", label_ctx.lower().rstrip(" ,:;("))
            for rx, gen in LABELS:
                if rx.search(low):
                    g_prev = gen
                    break
            self.prev = g_prev
            out.append(before)
            if key:
                out.append(ph(key, m.group()))
                self.found.append(key)
            else:
                out.append(m.group())
            pos = m.end()
        out.append(seg[pos:])
        if not DOTS.search(seg):
            self.triggers(seg.lower())
        return "".join(out)

# --------------------------------------------------------------------------
# Reguli explicite (aplicate inaintea detectiei automate)
# --------------------------------------------------------------------------
CHK_13 = [
    ("prima înregistrare în sistemul integrat", "CHK_PRIMA"),
    ("actualizare informaţii cadastrale", "CHK_ACTUALIZARE"),
    ("înregistrarea în planul cadastral a imobilelor situate în regiunile", "CHK_DL115"),
    ("identificare a amplasamentului imobilului situat pe un alt UAT", "CHK_ALTUAT"),
    ("atribuire număr cadastral", "CHK_ATRIBUIRE"),
    ("dezlipire/alipire teren", "CHK_DEZALIP"),
    ("primă înregistrare UI", "CHK_PRIMAUI"),
    ("documentaţie pentru apartamentare", "CHK_APART"),
    ("dezlipire/alipire UI", "CHK_DEZALIPUI"),
    ("documentaţie pentru reapartamentare", "CHK_REAPART"),
    ("supraetajare", "CHK_SUPRAETAJ"),
    ("dezmembrămintelor", "CHK_DEZMEMBR"),
    ("reconstituirea cărţii funciare", "CHK_RECONST"),
    ("îndreptarea erorii materiale a persoanei autorizate", "CHK_EROAREPA"),
    ("recepţie lucrări de specialitate", "CHK_RECLUCR"),
    ("avizare tehnică expertize judiciare", "CHK_EXPERTIZA"),
    ("exproprieri", "CHK_EXPROP"),
    ("întocmire plan parcelar", "CHK_PLANPARC"),
    ("întocmire plan încadrare în tarla", "CHK_INCTARLA"),
    ("trasarea imobilului", "CHK_TRASARE"),
    ("alte tipuri de lucrări de specialitate", "CHK_ALTELE"),
    ("intabularea", "CHK_INTAB"),
    ("înscrierea provizorie", "CHK_PROVIZORIE"),
    ("notarea", "CHK_NOTARE"),
    ("radierea", "CHK_RADIERE"),
    ("îndreptarea erorii materiale", "CHK_EROARE"),
    ("prin poştă", "CHK_POSTA"),
    ("la sediul", "CHK_SEDIU"),
    ("prin fax", "CHK_FAX"),
    ("prin e-mail", "CHK_EMAIL"),
    ("online în sistemul integrat", "CHK_ONLINE"),
    ("imobilul situat în: UAT", "CHK_IMOBIL"),
    ("zona identificată prin", "CHK_ZONA"),
]
CHK_14 = [
    ("am indicat persoanei autorizate limitele", "CHK_D1"),
    ("solicit înscrierea în evidenţele de cadastru", "CHK_D2"),
    ("poziţionarea incertă", "CHK_D3"),
    ("am adus la cunoştinţa tuturor proprietarilor", "CHK_D4"),
    ("litigiu", "CHK_D5"),
]
CHK_22 = [
    ("nu este înscris în cartea funciară şi documentaţia", "CHK_M1"),
    ("renunţ la operaţiunea de dezlipire/alipire", "CHK_M2"),
    ("suprapuneri reale", "CHK_M3"),
    ("sentinţă definitivă", "CHK_M4"),
    ("alt UAT ca urmare", "CHK_M5"),
]
CHK_23 = [("emiterea avizului de începere", "CHK_AVIZ"), ("recepţia tehnică a lucrării", "CHK_RECEPTIE")]
CHK_30 = [("lista titlurilor de proprietate", "CHK_I1"), ("extrasul de plan cadastral", "CHK_I2"),
          ("planuri parcelare existente", "CHK_I3"), ("copiile conforme", "CHK_I4"),
          ("lista cuprinzând numărul de carte funciară", "CHK_I5")]
CHK = {12: CHK_13, 13: CHK_13, 14: CHK_14, 22: CHK_22, 23: CHK_23, 30: CHK_30}
CHK_COMUNICARE = [("prin poştă", "CHK_POSTA"), ("la sediul BCPI", "CHK_SEDIU"), ("prin fax", "CHK_FAX"),
                  ("prin e-mail", "CHK_EMAIL"), ("online în sistemul integrat", "CHK_ONLINE")]

def apply_checkboxes(nr, text):
    rules = CHK.get(nr, []) + CHK_COMUNICARE
    if "[ ]" not in text:
        return text
    for frag, key in rules:
        if frag.lower() in text.lower():
            return text.replace("[ ]", ph(key, "[ ]"), 1)
    return text

# Inlocuiri explicite: (anexa, regex, inlocuire)
EXPLICIT = [
    # Anexa 13 - cerere de receptie si inscriere
    (13, r"^(\[ \] intabularea\*\) )\.+;", r"\1" + ph("INTAB_OBIECT", ".......") + ";"),
    (13, r"^(Imobilul: UAT )\.+", r"\1" + ph("UAT", ".........")),
    (13, r"(anexat în original sau în copie legalizată: )\.+", r"\1" + ph("ACT_PROP", "..........................")),
    (13, r"^(\.{60,})$", ph("P1_NUMEC", r"\1")),
    (13, r"^(\[ \] prin e-mail )\.+", r"\1" + ph("P1_EMAIL", ".................")),
    (5, r"(anexat în original sau în copie legalizată: )\.+", r"\1" + ph("ACT_PROP", "..............")),
    (5, r"^(\.{60,})$", ph("P1_NUMEC", r"\1")),
    (22, r"(anexat în original sau în copie legalizată: )\.+", r"\1" + ph("ACT_PROP", "......")),
    (22, r"^(ANULAREA NUMĂRULUI CADASTRAL: )\.+", r"\1" + ph("NRCAD", "........")),
    (22, r"(Cartea Funciară nr\. )\.+", r"\1" + ph("NRCF", ".....")),
    # Anexa 14 - declaratie
    (14, r"(al imobilului situat în )\.+", r"\1" + ph("ADRESA", "..........")),
    (14, r"(rezultate din măsurători de )\.+( mp)", r"\1" + ph("SMAS", "..................") + r"\2"),
    (14, r"(persoanei autorizate )\.+(/autorizaţie categoria/clasa )\.+( seria )\.+( nr\. )\.+",
         r"\1" + ph("E_NUMEC", "......................") + r"\2" + ph("E_CAT", "..............") +
         r"\3" + ph("E_AUT_SERIA", "..................") + r"\4" + ph("E_AUT_NR", "..................")),
    (14, r"(dreptului de proprietate )\.+(, în vederea)", r"\1" + ph("ACT_PROP", "..................") + r"\2"),
    (14, r"(cu imobilul ID )\.+(, nr\. dosar )\.+(, instanţa )\.+( obiect )\.+",
         r"\1" + ph("LIT_ID", "...................") + r"\2" + ph("LIT_DOSAR", "........................") +
         r"\3" + ph("LIT_INSTANTA", "..................") + r"\4" + ph("LIT_OBIECT", "..................")),
    # Anexa 15 - memoriu tehnic
    (15, r"^(1\. Adresa imobilului/zona în care se desfăşoară lucrarea) \.+", r"\1: " + ph("ADRESA", "..................")),
    (15, r"^(2\. Tipul lucrării:) \.+", r"\1 " + ph("TIP_LUCRARE", "..................")),
    (15, r"^(3\. Suprafaţa planului supus recepţiei) \.+", r"\1: " + ph("SMAS_MP", "..................")),
    (15, r"^(4\. Scurtă prezentare.*?istoricului) \.+", r"\1: " + ph("SCOP", "..................")),
    (15, r"^(5\. Descrierea generală.*?etc\.) \.+", r"\1: " + ph("DOCUMENTARE", ".....................")),
    (15, r"^(6\. Operaţiuni de specialitate realizate) \.+", r"\1: " + ph("OPERATIUNI", "..................")),
    # Anexa 19 - foaia colectiva (antet)
    (19, r"^(Carte Funciară Colectivă Nr\. )\.+(/UAT )\.+", r"\1" + ph("NRCF_COLECTIVA", ".......") + r"\2" + ph("UAT", "........")),
    (19, r"^(Adresa construcţiei: )\.+", r"\1" + ph("ADRESA", "...........")),
    (19, r"^(Nr\. cadastral: )\.+", r"\1" + ph("NRCAD", "........")),
    (19, r"^(Suprafaţă teren: )\.+", r"\1" + ph("SMAS_MP", "......")),
    # Anexa 20 - proces-verbal de vecinatate
    (20, r"^(Încheiat la data )\.+", r"\1" + ph("DATA", "..........")),
    (20, r"(încheiat în )\.+( exemplare)", r"\1" + ph("PV_EXEMPLARE", "....") + r"\2"),
    (20, r"(şi conţine )\.+( pagini)", r"\1" + ph("PV_PAGINI", "...") + r"\2"),
    (20, r"(denumirea locului)$", r"\1 " + ph("DENUMIRE_LOC", "")),
    # Anexa 23 - cerere CNC
    (23, r"^(\.{10,})$", None),  # tratate prin pozitie (vezi POZ_23)
    # Anexa 26-29 - formulare primarie
    (27, r"(vol\. )\.+( fila )\.+", r"\1" + ph("RA_VOL", "....") + r"\2" + ph("RA_FILA", ".....")),
    (27, r"(Certificăm că terenul în suprafaţă de )\.+( mp)", r"\1" + ph("SMAS", "....") + r"\2"),
    (27, r"(are categoria de folosinţă )\.+", r"\1" + ph("CATEG_TEXT", "..........")),
    (27, r"(în favoarea numitului )\.+", r"\1" + ph("P1_NUMEC", ".........")),
    (28, r"(edificată în anul )\.+", r"\1" + ph("C1_AN", "......")),
    (29, r"(edificată în anul )\.+", r"\1" + ph("C1_AN", ".....")),
    (28, r"^(– destinaţia )\.+", r"\1" + ph("C1_DEST", "...............")),
    (29, r"^(– destinaţia )\.+", r"\1" + ph("C1_DEST", "...........")),
    (28, r"^(– număr unităţi individuale )\.+", r"\1" + ph("C1_NRUI", "..........")),
    (29, r"^(– număr unităţi individuale )\.+", r"\1" + ph("C1_NRUI", "........")),
    (28, r"^(– regimul de înălţime )\.+", r"\1" + ph("C1_REGIM", "................")),
    (29, r"^(– regimul de înălţime )\.+", r"\1" + ph("C1_REGIM", "..............")),
    (28, r"^(– suprafaţa construită la sol )\.+", r"\1" + ph("C1_SC_MP", "........")),
    (29, r"^(– suprafaţa construită la sol )\.+", r"\1" + ph("C1_SC_MP", "......")),
    (28, r"^(– suprafaţa construită desfăşurată )\.+", r"\1" + ph("C1_SD_MP", ".........")),
    (29, r"^(– suprafaţa construită desfăşurată )\.+", r"\1" + ph("C1_SD_MP", ".........")),
    # Anexa 30 - cerere plan parcelar
    (30, r"(planului parcelar al tarlalei nr\. )\.+(/UAT )\.+", r"\1" + ph("TARLA", ".......") + r"\2" + ph("UAT", "...........")),
    (30, r"(posesoare a certificatului de autorizare )\.+(, eliberat de )\.+( la data )\.+(, conform contractului )\.+",
         r"\1" + ph("E_AUT", "........") + r"\2" + ph("E_AUT_EMITENT", "..........") + r"\3" +
         ph("E_AUT_DATA", "........") + r"\4" + ph("E_CONTRACT", "...........")),
    # Anexa 34 - registru proprietari
    (34, r"(Subsemnata/Subsemnatul/Subscrisa )\.+", r"\1" + ph("P1_NUMEC", "...........")),
]

# Campuri asezate dupa pozitia paragrafului (anexa 23: II, III, IV)
POZ_23 = {"II. TIPUL LUCRĂRII:": "TIP_LUCRARE", "III. EXECUTANT:": "E_NUMEC_AUT",
          "IV. BENEFICIAR:": "P1_NUMEC"}

def explicit(nr, text):
    for n, rx, rep in EXPLICIT:
        if n == nr and rep is not None:
            text = re.sub(rx, rep, text)
    return text

# --------------------------------------------------------------------------
# Sabloane speciale pentru tabele
# --------------------------------------------------------------------------
ADRESA_IMOBIL_CELLS = ["LOCALITATE", "STRADA_TARLA", "NR_PARCELA", "BLOC", "SCARA", "ETAJ", "AP", "NRCF_NRCAD"]

def fill_tables(nr, block, ti):
    """Completeaza celulele tabelelor cu campuri/randuri repetabile."""
    rows = block["rows"]
    first = rows[0][0]["text"] if rows and rows[0] else ""
    # tabelul "Adresă imobil" (anexele 12, 13, 23)
    if first.startswith("Adresă imobil") and len(rows) >= 3:
        cells = rows[2]
        for i, c in enumerate(cells[:8]):
            c["text"] = ph(ADRESA_IMOBIL_CELLS[i], "")
        return block
    # anexa 12: persoana autorizata solicitanta
    if nr == 12 and first.startswith("Nume (denumire PJA)") and len(rows) >= 3:
        keys = ["E_NUME", "E_PRENUME", "E_CNP", "E_CAT", "E_AUT_SERIA", "E_AUT_NR"]
        for i, c in enumerate(rows[2][:6]):
            c["text"] = ph(keys[i], c["text"] if keys[i] == "E_CNP" else "")
        return block
    # anexele 1, 5-8, 10: modalitatea de comunicare (bife in celule)
    for r in rows:
        for c in r:
            if "[ ]" in c["text"]:
                c["text"] = apply_checkboxes(nr, c["text"])
    # anexa 20: tabele cu randuri repetabile
    if nr == 20:
        if ti == 0:   # proprietarii imobilului
            hdr = rows[0]
            tpl = [ph("PROP.NUME", ""), ph("PROP.PRENUME", ""), ph("PROP.CALITATE", "")] + \
                  [ph("PROP.CNP#%d" % (i + 1), "") for i in range(13)]
            block["rows"] = [hdr, [dict(text=t, colspan=1, rowspan=1, center=True, noborder=False) for t in tpl]]
        elif ti == 1:  # proprietarii vecini
            hdr = rows[0]
            tpl = [ph("VEC.NUME", ""), ph("VEC.PRENUME", ""), ph("VEC.CALITATE", "")] + \
                  [ph("VEC.CNP#%d" % (i + 1), "") for i in range(13)] + [ph("VEC.CUI", ""), ph("VEC.OBS", "")]
            block["rows"] = [hdr, [dict(text=t, colspan=1, rowspan=1, center=True, noborder=False) for t in tpl]]
        elif ti == 2:  # laturile
            hdr = rows[0]
            tpl = [ph("LAT.LATURA", ""), ph("LAT.LUNG", ""), ph("LAT.MAT", ""), ph("LAT.VECIN", ""),
                   ph("LAT.ID", ""), ph("LAT.OBS", "")]
            block["rows"] = [hdr, [dict(text=t, colspan=1, rowspan=1, center=True, noborder=False) for t in tpl]]
        elif ti == 3:  # semnaturile
            left, right = rows[0][0], rows[0][1]
            left["text"] = left["text"] + "\n" + ph("SEMN_PROP", "")
            right["text"] = (right["text"]
                             .replace("Latura ............ (semnătura).......",
                                      ph("SEMN_VECINI", "Latura ............ (semnătura)......."))
                             .replace("Nume prenume/denumire", "Nume prenume/denumire: " + ph("E_NUMEC", ""))
                             .replace("Seria și nr. certificatului de autorizare",
                                      "Seria și nr. certificatului de autorizare: " + ph("E_AUT", ""))
                             .replace("\nCategoria\n", "\nCategoria: " + ph("E_CAT", "") + "\n"))
        return block
    return block

def custom_table_19():
    hdr = ["Nr. tronson", "Scară", "Nivel", "Nr. apartament/ spațiu cu altă destinație",
           "Cod Unitate individuală\n(UI)*", "Nr. CF individuală*", "Suprafața utilă UI – Suap - (mp)",
           "Cotă parte indiviză din părțile comune aferente UI\n(Cpc)", "Cotă parte teren aferent\nUI\n(Cpt)", "Obs."]
    keys = ["TRONSON", "SCARA", "NIVEL", "NRAP", "COD", "CF", "SU", "CPC", "CPT", "OBS"]
    C = lambda t, cs=1: dict(text=t, colspan=cs, rowspan=1, center=True, noborder=False)
    rows = [[C(h) for h in hdr], [C(str(i + 1)) for i in range(10)],
            [C(ph("UI." + k, "")) for k in keys],
            [C("TOTAL", 6), C(ph("TOT_SU", "")), C(ph("TOT_CPC", "")), C(ph("TOT_CPT", "")), C("")],
            [C("Descriere părți comune indivize", 3),
             C(ph("PARTI_COMUNE", "(Ex. intrare, casa scării, subsol tehnic, spălătorie, uscătorie, ghenă gunoi, "
                  "pod, acoperiș, terenul construit și neconstruit, etc.)"), 7)],
            [C("Întocmit\n" + ph("E_NUMEC", "") + "\n(semnătură)\nData " + ph("DATA", "..................."), 6),
             C("Recepționat\n(semnătură)", 4)]]
    return {"t": "table", "rows": rows, "widths": [1.3, 1.1, 1.8, 2.6, 2.8, 2.0, 2.0, 2.2, 2.4, 5.6]}

def custom_table_32():
    hdr = ["Nr. Tarla", "Nr. crt.\n(ordinea în tarla)", "Nr. cadastral inițial", "Nr. carte funciară",
           "Nr. parcelă/(nr. topografic)", "Nr. Cad/Nr. CF**)", "Suprafața parcelă din act\n(mp)",
           "Categoria de folosință din act", "Suprafața totală imobil din act\n(mp)",
           "Suprafața totală imobil determinată\n(mp)", "Numele și prenumele proprietarului",
           "Act de proprietate\n(nr./data)\n*****)", "Observații"]
    keys = ["TARLA", "NRCRT", "NRCADI", "NRCF", "PARCELA", "NRCADF", "SPACT", "CATEG", "STACT", "STMAS",
            "PROPRIETAR", "ACT", "OBS"]
    C = lambda t, cs=1: dict(text=t, colspan=cs, rowspan=1, center=True, noborder=False)
    rows = [[C(h) for h in hdr], [C(str(i)) for i in range(13)],
            [C(ph("TP." + k, "")) for k in keys],
            [C(""), C("TOTAL SUPRAF. TARLA", 5), C(ph("TOT_SPACT", "")), C(""), C(ph("TOT_STACT", "")),
             C(ph("TOT_STMAS", "")), C("", 3)]]
    return {"t": "table", "rows": rows, "widths": [1.2, 1.4, 1.8, 1.8, 2.0, 1.8, 1.8, 1.8, 1.9, 1.9, 3.2, 2.4, 2.4]}

def inventar_blocks():
    """Sablon suplimentar: inventarul de coordonate (nu face parte din ordin)."""
    C = lambda t, cs=1: dict(text=t, colspan=cs, rowspan=1, center=True, noborder=False)
    return [
        {"t": "p", "text": "INVENTAR DE COORDONATE", "align": "center", "lvl": 0},
        {"t": "p", "text": "Sistem de proiecţie Stereografic 1970", "align": "center", "lvl": 0},
        {"t": "p", "text": "Imobil: nr. cadastral " + ph("NRCAD", "........") + ", CF nr. " + ph("NRCF", "........") +
              ", UAT " + ph("UAT", "........"), "align": "left", "lvl": 0},
        {"t": "p", "text": "Adresa: " + ph("ADRESA", "........"), "align": "left", "lvl": 0},
        {"t": "table", "rows": [[C("Nr. pct."), C("X [m]"), C("Y [m]"), C("Latura"), C("Lungime [m]")],
                                [C(ph("PCT.NR", "")), C(ph("PCT.X", "")), C(ph("PCT.Y", "")),
                                 C(ph("PCT.LAT", "")), C(ph("PCT.DIST", ""))]]},
        {"t": "p", "text": "S(" + ph("NRCAD", "") + ") = " + ph("SMAS", "......") + " mp      P = " +
              ph("PERIM", "......") + " m", "align": "left", "lvl": 0},
        {"t": "p", "text": "Întocmit,", "align": "right", "lvl": 0},
        {"t": "p", "text": ph("E_NUMEC", "........"), "align": "right", "lvl": 0},
        {"t": "p", "text": "Data: " + ph("DATA", "........"), "align": "right", "lvl": 0},
    ]

def memoriu_extra():
    """Completare la memoriul tehnic: constructii si inventar de coordonate."""
    C = lambda t, cs=1: dict(text=t, colspan=cs, rowspan=1, center=True, noborder=False)
    return [
        {"t": "p", "text": ph("CONSTR_TEXT", ""), "align": "left", "lvl": 1},
        {"t": "p", "text": "Inventarul de coordonate al punctelor de pe conturul imobilului (sistem Stereografic 1970):",
         "align": "left", "lvl": 1},
        {"t": "table", "rows": [[C("Nr. pct."), C("X [m]"), C("Y [m]"), C("Latura"), C("Lungime [m]")],
                                [C(ph("PCT.NR", "")), C(ph("PCT.X", "")), C(ph("PCT.Y", "")),
                                 C(ph("PCT.LAT", "")), C(ph("PCT.DIST", ""))]]},
        {"t": "p", "text": "Suprafaţa măsurată S = " + ph("SMAS", "......") + " mp, perimetrul P = " +
              ph("PERIM", "......") + " m.", "align": "left", "lvl": 1},
    ]

# --------------------------------------------------------------------------
# Prelucrarea unei anexe
# --------------------------------------------------------------------------
class _Off:
    """Motor inactiv: pastreaza textul (anexele completate de OCPI/CNC)."""
    found = []
    def run(self, text):
        return text

def process(nr, blocks):
    eng = _Off() if nr in AUTO_OFF else Engine(nr, exec_subject=(nr == 23))
    out, ti = [], 0
    for b in blocks:
        if b["t"] == "p":
            t = b["text"]
            if nr in (16, 17) and t.startswith("*) Anexa nr."):
                continue
            t = explicit(nr, t)
            t = apply_checkboxes(nr, t)
            if nr == 23 and re.fullmatch(r"\.{10,}", t) and out and out[-1]["t"] == "p":
                key = POZ_23.get(out[-1]["text"].strip())
                if key:
                    t = ph(key, t)
            t = eng.run(t)
            out.append(dict(b, text=t))
        else:
            if nr == 19:
                out.append(custom_table_19())
            elif nr == 32 and ti == 0:
                out.append(custom_table_32())
            else:
                tb = json.loads(json.dumps(b))
                tb = fill_tables(nr, tb, ti)
                for r in tb["rows"]:
                    for c in r:
                        if "{{" not in c["text"]:
                            c["text"] = eng.run(c["text"])
                out.append(tb)
            ti += 1
        if nr == 15 and b["t"] == "p" and b["text"].startswith("(prezentarea succintă"):
            out.extend(memoriu_extra())
    return out

# --------------------------------------------------------------------------
# Scrierea documentelor Word
# --------------------------------------------------------------------------
def set_cell_border(cell, show=True):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single" if show else "nil")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "000000")
        borders.append(el)
    tcPr.append(borders)

def add_runs(par, text, size=None, bold=False, italic=False):
    """Adauga textul; fiecare camp {{...}} intr-un run separat (completare sigura)."""
    lines = text.split("\n")
    for li, line in enumerate(lines):
        pos = 0
        for m in PH.finditer(line):
            if m.start() > pos:
                r = par.add_run(line[pos:m.start()])
                r.bold, r.italic = bold, italic
                if size: r.font.size = Pt(size)
            r = par.add_run(m.group())
            r.bold, r.italic = bold, italic
            if size: r.font.size = Pt(size)
            pos = m.end()
        if pos < len(line) or not line:
            r = par.add_run(line[pos:])
            r.bold, r.italic = bold, italic
            if size: r.font.size = Pt(size)
        if li < len(lines) - 1:
            par.add_run().add_break()

def html_grid(rows):
    """Transforma randurile cu colspan/rowspan intr-o grila (r, c) -> celula."""
    grid, placed = {}, []
    for r, row in enumerate(rows):
        c = 0
        for cell in row:
            while (r, c) in grid:
                c += 1
            for dr in range(cell["rowspan"]):
                for dc in range(cell["colspan"]):
                    grid[(r + dr, c + dc)] = cell
            placed.append((r, c, cell))
            c += cell["colspan"]
    nrows = max(k[0] for k in grid) + 1
    ncols = max(k[1] for k in grid) + 1
    return placed, nrows, ncols

def column_widths(block, nc, usable):
    """Latimi de coloana (cm): proportiile date de tabel ("widths") sau coloane
    inguste pentru cifrele CNP (13 celule)."""
    if block.get("widths") and len(block["widths"]) == nc:
        tot = float(sum(block["widths"]))
        return [usable * w / tot for w in block["widths"]]
    digits = set()
    for r, row in enumerate(block["rows"]):
        c = 0
        for cell in row:
            if cell["colspan"] == 13 and cell["text"].strip().upper().startswith("CNP"):
                digits.update(range(c, c + 13))
            c += cell["colspan"]
    if not digits:
        return None
    dw = 0.42
    rest = (usable - dw * len(digits)) / max(1, nc - len(digits))
    return [dw if i in digits else rest for i in range(nc)]

def write_table(doc, block, small, usable=17.5):
    placed, nr, nc = html_grid(block["rows"])
    t = doc.add_table(rows=nr, cols=nc)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = column_widths(block, nc, usable)
    if widths:
        t.autofit = False
        for i, col in enumerate(t.columns):
            col.width = Cm(widths[i])
        for row in t.rows:
            for i, cell in enumerate(row.cells):
                cell.width = Cm(widths[i])
    for r, c, cell in placed:
        a = t.cell(r, c)
        b = t.cell(min(r + cell["rowspan"] - 1, nr - 1), min(c + cell["colspan"] - 1, nc - 1))
        if (a._tc is not b._tc):
            a = a.merge(b)
        p = a.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if cell.get("center", True):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, cell["text"], size=small)
        if cell.get("noborder"):
            set_cell_border(a, False)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)

def write_doc(nr, blocks, path, title_extra=None):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    if nr in LANDSCAPE:
        sec.orientation = WD_ORIENT.LANDSCAPE
        sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
    sec.left_margin, sec.right_margin = Cm(2.0), Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)
    st = doc.styles["Normal"]
    st.font.name = "Times New Roman"
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(3)
    # antetul "Anexa nr. X la regulament"
    if nr:
        hp = doc.add_paragraph()
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        add_runs(hp, "Anexa nr. %d la regulament" % nr, size=10, italic=True)
    for b in blocks:
        if b["t"] == "table":
            small = 8 if nr in (12, 19, 20, 32) else 10
            write_table(doc, b, small, 26.2 if nr in LANDSCAPE else 17.5)
            continue
        t = b["text"]
        if nr and t == "la regulament":
            continue
        p = doc.add_paragraph()
        al = b.get("align", "left")
        p.alignment = {"center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}.get(
            al, WD_ALIGN_PARAGRAPH.JUSTIFY if len(t) > 110 else WD_ALIGN_PARAGRAPH.LEFT)
        if b.get("lvl"):
            p.paragraph_format.left_indent = Cm(0.6 * b["lvl"])
        plain = PH.sub("", t)
        is_title = al == "center" and plain.isupper() and len(plain) > 4
        is_note = bool(re.match(r"^\*+\)", plain)) or plain.startswith("Notă") or plain.startswith("NOTĂ")
        size = 9 if (is_note or plain.startswith("Prin semnarea") or plain.startswith("Datele cu caracter")
                     or plain.startswith("Oficiul de cadastru şi publicitate imobiliară, instituţie")
                     or plain.startswith("În exercitarea drepturilor")) else None
        add_runs(p, t, size=13 if is_title else size, bold=is_title)
    doc.core_properties.title = "Anexa nr. %d - %s" % (nr, TITLURI.get(nr, "")) if nr else (title_extra or "")
    doc.core_properties.author = "TopoKit"
    doc.save(path)

def main():
    raport = "--raport" in sys.argv
    data = json.load(open(SRC, encoding="utf-8"))
    OUT.mkdir(exist_ok=True)
    catalog = []
    for k in sorted(data, key=int):
        nr = int(k)
        if nr in GRAFICE:
            continue
        blocks = process(nr, data[k])
        fn = "Anexa_%02d.docx" % nr
        write_doc(nr, blocks, OUT / fn)
        keys = sorted(set(m.group(1).split("|")[0] for b in blocks for m in
                          re.finditer(r"\{\{([^{}]*)\}\}", json.dumps(b, ensure_ascii=False))))
        catalog.append({"nr": nr, "titlu": TITLURI.get(nr, ""), "fisier": fn, "campuri": keys})
        if raport:
            print("==== Anexa %d - %s" % (nr, TITLURI.get(nr, "")))
            for b in blocks:
                txt = b["text"] if b["t"] == "p" else " | ".join(c["text"] for r in b["rows"] for c in r)
                if "{{" in txt:
                    print("   ", re.sub(r"\|[^{}]*\}\}", "}}", txt)[:400])
    write_doc(0, inventar_blocks(), OUT / "Inventar_coordonate.docx", "Inventar de coordonate")
    catalog.append({"nr": 0, "titlu": "Inventar de coordonate", "fisier": "Inventar_coordonate.docx",
                    "campuri": ["NRCAD", "NRCF", "UAT", "ADRESA", "SMAS", "PERIM", "E_NUMEC", "DATA", "PCT.*"]})
    json.dump(catalog, open(OUT / "catalog.json", "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print("%d sabloane scrise in %s" % (len(catalog), OUT))

if __name__ == "__main__":
    main()
